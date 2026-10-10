"""Exportadores CSV / XLSX / PDF.

Patron Strategy + Registry. ``register_exporter`` y ``get_exporter``
permiten agregar nuevos formatos sin modificar el cliente.
"""

from __future__ import annotations

import csv
import io
from abc import ABC, abstractmethod

import openpyxl
from django.http import HttpResponse
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .datasets import TabularData


class BaseExporter(ABC):
    """Contrato base de los exportadores."""

    format_key: str = ""
    content_type: str = ""
    extension: str = ""

    @abstractmethod
    def render(self, data: TabularData) -> bytes:
        """Serializa ``data`` a bytes."""

    def to_response(self, data: TabularData) -> HttpResponse:
        """Devuelve un ``HttpResponse`` con los headers correctos."""
        body = self.render(data)
        response = HttpResponse(
            body, content_type=self.content_type,
        )
        response["Content-Disposition"] = (
            f'attachment; filename="{data.filename}.{self.extension}"'
        )
        return response


class CsvExporter(BaseExporter):
    """Exportador CSV con BOM para Excel."""

    format_key = "csv"
    content_type = "text/csv; charset=utf-8"
    extension = "csv"

    def render(self, data: TabularData) -> bytes:
        """Genera un CSV en memoria con BOM."""
        buffer = io.StringIO()
        buffer.write("﻿")
        writer = csv.writer(buffer)
        writer.writerow(data.headers)
        writer.writerows(data.rows)
        return buffer.getvalue().encode("utf-8")


class XlsxExporter(BaseExporter):
    """Exportador Excel via openpyxl."""

    format_key = "xlsx"
    content_type = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    extension = "xlsx"

    def render(self, data: TabularData) -> bytes:
        """Genera un workbook en memoria."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = data.title[:31]
        ws.append(data.headers)
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        for row in data.rows:
            ws.append(row)
        for col_idx, _ in enumerate(data.headers, start=1):
            ws.column_dimensions[
                openpyxl.utils.get_column_letter(col_idx)
            ].width = 18
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()


class PdfExporter(BaseExporter):
    """Exportador PDF via reportlab."""

    format_key = "pdf"
    content_type = "application/pdf"
    extension = "pdf"

    def render(self, data: TabularData) -> bytes:
        """Genera un PDF simple en memoria."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        styles = getSampleStyleSheet()
        story: list = [Paragraph(data.title, styles["Title"]), Spacer(1, 12)]
        table = Table([data.headers, *data.rows])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d6efd")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
        ]))
        story.append(table)
        doc.build(story)
        return buffer.getvalue()


EXPORTERS: dict[str, type[BaseExporter]] = {}


def register_exporter(cls: type[BaseExporter]) -> type[BaseExporter]:
    """Decorador que anade el exporter al registro."""
    EXPORTERS[cls.format_key] = cls
    return cls


def get_exporter(fmt: str) -> BaseExporter:
    """Devuelve el exporter para ``fmt`` (``csv`` por defecto)."""
    cls = EXPORTERS.get(fmt, CsvExporter)
    return cls()


for exporter_cls in (CsvExporter, XlsxExporter, PdfExporter):
    register_exporter(exporter_cls)


__all__ = [
    "BaseExporter",
    "CsvExporter",
    "XlsxExporter",
    "PdfExporter",
    "EXPORTERS",
    "register_exporter",
    "get_exporter",
]
