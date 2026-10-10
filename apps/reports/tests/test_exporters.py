"""Tests de exportadores."""

from __future__ import annotations

import csv

import openpyxl

from apps.reports.datasets import TabularData
from apps.reports.exporters import (
    CsvExporter,
    PdfExporter,
    XlsxExporter,
    get_exporter,
)


def sample_data() -> TabularData:
    """Crea un dataset de prueba."""
    return TabularData(
        title="Demo",
        headers=["col_a", "col_b"],
        rows=[["1", "x"], ["2", "z"]],
        filename="sample",
    )


def test_csv_exporter_writes_bom_and_rows() -> None:
    """El CSV incluye BOM y los datos correctos."""
    body = CsvExporter().render(sample_data())
    assert body.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(body[3:].decode("utf-8").splitlines()))
    assert rows[0] == ["col_a", "col_b"]
    assert rows[1:] == [["1", "x"], ["2", "z"]]


def test_xlsx_exporter_returns_workbook() -> None:
    """El XLSX genera un workbook con la primera fila en negrita."""
    body = XlsxExporter().render(sample_data())
    wb = openpyxl.load_workbook(__import__("io").BytesIO(body))
    ws = wb.active
    assert ws["A1"].value == "col_a"
    assert ws["A1"].font.bold is True
    assert ws.cell(row=2, column=1).value == "1"


def test_pdf_exporter_starts_with_pdf_header() -> None:
    """El PDF empieza con la firma ``%PDF``."""
    body = PdfExporter().render(sample_data())
    assert body.startswith(b"%PDF")


def test_get_exporter_returns_default_for_unknown() -> None:
    """``get_exporter`` devuelve ``CsvExporter`` si el formato no existe."""
    exporter = get_exporter("nope")
    assert isinstance(exporter, CsvExporter)


def test_get_exporter_returns_correct_class() -> None:
    """``get_exporter`` devuelve la clase adecuada."""
    assert isinstance(get_exporter("xlsx"), XlsxExporter)
    assert isinstance(get_exporter("pdf"), PdfExporter)
